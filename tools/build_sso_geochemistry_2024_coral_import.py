#!/usr/bin/env python3
"""Build the CORAL import package for the 2024 SSO groundwater geochemistry.

Sources (team Google Drive mirror):

* Chakraborty lab ion chromatography, 2024 pump test (17 samples x 3 technical
  replicates; chloride, sulfate, nitrate as the NO3- ion; 1:50 dilution).
* Chakraborty lab TOC/TN, 2024 pump test (17 samples x 3 technical replicates).
* Chakraborty lab ion chromatography, 2024 pilot time series (38 samples).
* Stahl/Hunt lab nitrate isotopes (Wankel lab denitrifier method), 2024 pump
  test (17 samples x 2 analytical replicates): d15N, d18O, nitrate uM.

Every sample maps to an existing CORAL Sample by name, except the 200 L
isotope sample of the M5-SZ3 pumping series, which the field workbook recorded
under the 100 L name; it is added as a static Sample with its own Sampling
process.

Outputs: ``coral_import/sso_geochemistry_20261009/`` with ``json/``,
``static/``, ``process/``, ``check/``, ``reports/``, ``files_to_import.txt``,
``import_to_coral.py`` and ``README.md``.
"""

from __future__ import annotations

import csv
import json
import os
import subprocess
import sys
from collections import OrderedDict
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parents[1]
DRIVE = Path("/mnt/net/dipa.jmcnet/data/backup/google-drive/enigma_team_drive_data_enigma_backup")
EXPORT = REPO / "sync-coral-to-berdl/exports/sync-20261009-134304/coral_export/static_tsv"
ONTOLOGIES = Path("/home/coral/prod/data_import/ontologies")
RELEASE = "20261009"
OUT = REPO / f"coral_import/sso_geochemistry_{RELEASE}"
JSON_OUT, STATIC_OUT, PROCESS_OUT, CHECK_OUT, REPORT_OUT = (OUT / d for d in ("json", "static", "process", "check", "reports"))

IC_PUMP = DRIVE / "2024_SSO_Pump Test/Chakraborty_IC／TOC/20240412_SSO_pumptest_IC_full_w_rawdata.xlsx"
TOC_PUMP = DRIVE / "2024_SSO_Pump Test/Chakraborty_IC／TOC/20240416_TOC_TN_SSO_pumptest_corrected.xlsx"
IC_TS = DRIVE / "2024_SSO_Pilot Time Series/Chakraborty_Anions/20241220_IC_results_SSO_timeseries.xlsx"
ISO_PUMP = DRIVE / "2024_SSO_Pump Test/Stahl_Isotopes/20240416_TOC_TN_IC_ISOTOPES_SSO_pumptest.xlsx"
DRIVE_LABEL = "ENIGMA team Google Drive"

CHECKGENERIC_CLASSES = Path("/h/jmc/src/java/classes")
JACKSON_JARS = [
    Path("/h/jmc/.m2/repository/com/fasterxml/jackson/core/jackson-annotations/2.5.1/jackson-annotations-2.5.1.jar"),
    Path("/h/jmc/.m2/repository/com/fasterxml/jackson/core/jackson-core/2.5.1/jackson-core-2.5.1.jar"),
    Path("/h/jmc/.m2/repository/com/fasterxml/jackson/core/jackson-databind/2.5.1/jackson-databind-2.5.1.jar"),
]

# ontology ids (verified against the production OBO files at build time)
T = {
    "chemical measurement": "DA:0000005",
    "data variables type": "ME:0000293",
    "instrument": "ME:0000012",
    "method": "ME:0000007",
    "comment": "ME:0000011",
    "link": "ME:0000203",
    "personnel": "ENIGMA:0000029",
    "environmental sample": "ME:0000100",
    "environmental sample ID": "ME:0000102",
    "molecule": "ME:0000027",
    "molecule from list": "ME:0000381",
    "replicate series": "ME:0000005",
    "experimental context": "ME:0000002",
    "technical replicate": "ME:0000287",
    "concentration": "ME:0000129",
    "below": "ME:0000363",
    "relative": "ME:0000020",
    "organic": "ME:0000346",
    "isotope ratio": "ME:0000403",
    "isotope": "ME:0000408",
    "ppm": "UO:0000169",
    "per mil": "UO:0000168",
    "uM": "UO:0000064",
    "count unit": "UO:0000189",
    "chloride": "CHEBI:17996",
    "sulfate": "CHEBI:16189",
    "nitrate": "CHEBI:17632",
    "elemental carbon": "CHEBI:33415",
    "nitrogen atom": "CHEBI:25555",
    "nitrogen-15 atom": "CHEBI:36934",
    "oxygen-18 atom": "CHEBI:33815",
    "Romy Chakraborty": "ENIGMA:0000054",
    "Kristopher Hunt": "ENIGMA:0000164",
}
NAMES: dict[str, str] = {}
TS_DATE_CONFLICTS: list[str] = []


def load_ontology_names() -> None:
    for obo in sorted(ONTOLOGIES.glob("*.obo")):
        cur = None
        with obo.open(errors="ignore") as h:
            for line in h:
                line = line.strip()
                if line == "[Term]":
                    cur = None
                elif line.startswith("id: "):
                    cur = line[4:].strip()
                elif line.startswith("name: ") and cur:
                    NAMES.setdefault(cur, line[6:].strip())
                    cur = None
    missing = [k for k, v in T.items() if v not in NAMES]
    if missing:
        raise SystemExit(f"ontology ids not found: {missing}")


def term(key: str) -> dict:
    ref = T[key]
    return {"term_name": NAMES[ref], "oterm_ref": ref, "oterm_name": NAMES[ref]}


def oterm_value(key: str) -> dict:
    return {"scalar_type": "oterm_ref", "oterm_ref": T[key], "string_value": NAMES[T[key]]}


def ctx(type_key: str, value: dict) -> dict:
    return {"value_type": term(type_key), "value": value}


def string_value(text: str) -> dict:
    return {"scalar_type": "string", "string_value": text}


def boolean_value(flag: bool) -> dict:
    # scalar booleans are stored as 0/1 integers by the generics Value class
    return {"scalar_type": "boolean", "boolean_value": 1 if flag else 0}


# ----------------------------------------------------------------------------
# sources
# ----------------------------------------------------------------------------

def rows(path: Path, sheet: str | None = None) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it)]
    out = []
    for r in it:
        if not any(v not in (None, "") for v in r):
            continue
        out.append({header[i]: r[i] for i in range(min(len(header), len(r))) if header[i]})
    return out


PERIOD_SUFFIX = {"PRE": "1", "POST": "2", "0L": "", "100L": "100", "200L": "200"}
PUMP_DATES = {  # well-zone -> sampling date of the pump test (PRE/POST day)
    "U2-SZ1": "20240313", "L8-SZ1": "20240314", "M5-SZ1": "20240318", "U2-SZ2": "20240319",
    "L8-SZ2": "20240320", "M5-SZ2": "20240321", "M5-SZ3": "20240328",
}


def pump_sample(well: str, zone: str, period: str, assay: str) -> str:
    wz = f"{well}-{zone}"
    if period in ("0L", "100L", "200L"):
        return f"{wz}-20240329-{assay}{PERIOD_SUFFIX[period]}"
    return f"{wz}-{PUMP_DATES[wz]}-{assay}{PERIOD_SUFFIX[period]}"


def fnum(v) -> float | None:
    if v is None or v == "":
        return None
    return float(v)


# ----------------------------------------------------------------------------
# brick assembly
# ----------------------------------------------------------------------------

def sample_dim(names: list[str]) -> dict:
    return {
        "data_type": term("environmental sample"),
        "size": len(names),
        "typed_values": [
            {
                "value_type": term("environmental sample ID"),
                "values": {"scalar_type": "object_ref", "object_refs": list(names), "string_values": list(names)},
            }
        ],
    }


def molecule_dim(keys: list[str]) -> dict:
    return {
        "data_type": term("molecule"),
        "size": len(keys),
        "typed_values": [
            {
                "value_type": term("molecule from list"),
                "values": {"scalar_type": "oterm_ref", "oterm_refs": [T[k] for k in keys], "string_values": [NAMES[T[k]] for k in keys]},
            }
        ],
    }


def replicate_dim(n: int) -> dict:
    return {
        "data_type": term("replicate series"),
        "size": n,
        "typed_values": [
            {
                "value_type": term("replicate series"),
                "value_units": term("count unit"),
                "value_context": [ctx("experimental context", oterm_value("technical replicate"))],
                "values": {"scalar_type": "int", "int_values": list(range(1, n + 1))},
            }
        ],
    }


def float_var(type_key: str, unit_key: str, context: list[dict], values: list[float | None]) -> dict:
    return {"value_type": term(type_key), "value_units": term(unit_key), "value_context": context, "values": {"scalar_type": "float", "float_values": values}}


def bool_var(type_key: str, context: list[dict], values: list[bool | None]) -> dict:
    # the generics Values class stores booleans as 0/1 integers (null for missing)
    encoded = [None if v is None else (1 if v else 0) for v in values]
    return {"value_type": term(type_key), "value_context": context, "values": {"scalar_type": "boolean", "boolean_values": encoded}}


def below_ctx(extra: list[dict] | None = None) -> list[dict]:
    return [ctx("relative", string_value("detection limit"))] + (extra or [])


def brick(name: str, description: str, array_context: list[dict], dims: list[dict], variables: list[dict]) -> dict:
    marker = {"value_type": term("data variables type"), "value": {"scalar_type": "oterm_ref", "oterm_ref": T["chemical measurement"], "string_value": NAMES[T["chemical measurement"]]}}
    return {
        "name": f"{name}.hndarray",
        "description": description,
        "data_type": term("chemical measurement"),
        "array_context": [marker] + array_context,
        "n_dimensions": len(dims),
        "dim_context": dims,
        "typed_values": variables,
    }


IC_METHOD = (
    "Dionex Integrion HPIC with AS-AP autosampler, 30 mM KOH eluent at 0.25 mL/min, column and "
    "compartment at 30 C, suppressor on; mixed anion standards; samples diluted 1:50 in nanopure water; "
    "concentrations are reported as the anion (chloride, sulfate and nitrate as NO3-)."
)


def build_ic(samples: list[str], rep_table: dict[tuple[str, int], dict[str, float | None]], n_rep: int, name: str, desc: str, extra_context: list[dict]):
    molecules = ["chloride", "sulfate", "nitrate"]
    conc: list[float | None] = []
    below: list[bool | None] = []
    for s in samples:
        for m in molecules:
            for r in range(1, n_rep + 1):
                v = rep_table.get((s, r), {}).get(m)
                if v is None:
                    conc.append(None)
                    below.append(None)
                elif v <= 0:
                    conc.append(None)
                    below.append(True)
                else:
                    conc.append(round(v, 4))
                    below.append(False)
    dims = [sample_dim(samples), molecule_dim(molecules)]
    if n_rep > 1:
        dims.append(replicate_dim(n_rep))
    array_context = [
        ctx("instrument", string_value("Dionex Integrion HPIC")),
        ctx("method", string_value(IC_METHOD)),
        ctx("personnel", oterm_value("Romy Chakraborty")),
    ] + extra_context
    return brick(name, desc, array_context, dims, [
        float_var("concentration", "ppm", [], conc),
        bool_var("below", below_ctx(), below),
    ])


def build_pump_ic() -> tuple[dict, list[str]]:
    data = rows(IC_PUMP)
    table: dict[tuple[str, int], dict[str, float | None]] = {}
    samples: list[str] = []
    for r in data:
        s = pump_sample(r["Well"], r["Zone"], r["sampling_period"], "ICTOC")
        if s not in samples:
            samples.append(s)
        table[(s, int(r["technical_replicate"]))] = {"chloride": fnum(r["Chloride_ppm"]), "sulfate": fnum(r["Sulfate_ppm"]), "nitrate": fnum(r["Nitrate_ppm"])}
    extra = [ctx("comment", string_value(f"Source: {DRIVE_LABEL} 2024_SSO_Pump Test/Chakraborty_IC／TOC/20240412_SSO_pumptest_IC_full_w_rawdata.xlsx; triplicate injections of each sample"))]
    return build_ic(samples, table, 3, "chakraborty_ic_anions_sso_pump_test_2024", "Chakraborty Lab ion chromatography anions (chloride, sulfate, nitrate) for the 2024 SSO pump test groundwater samples", extra), samples


def build_ts_ic() -> tuple[dict, list[str]]:
    data = rows(IC_TS)
    table: dict[tuple[str, int], dict[str, float | None]] = {}
    samples: list[str] = []
    # The campaign sampled every port on 2024-09-09 (T1) and 2024-09-18 (T2).
    # The workbook's Date Sampled column disagrees with the T1/T2 label on six
    # rows (L8 swapped, L9-SZ1 and U2-SZ1 duplicated), so the timepoint label
    # is the mapping key and the disagreements are reported.
    timepoint_dates = {"T1": "20240909", "T2": "20240918"}
    date_conflicts = []
    for r in data:
        period = str(r["sampling_period"]).strip()
        date = timepoint_dates[period]
        if str(r["Date Sampled"]).strip() != date:
            date_conflicts.append(f"{r['SampleID']} (workbook date {r['Date Sampled']}, timepoint date {date})")
        s = f"{r['Well']}-{r['Zone']}-{date}-ICTOC"
        if s in samples:
            raise ValueError(f"duplicate time-series sample {s}")
        samples.append(s)
        table[(s, 1)] = {"chloride": fnum(r["Chloride_ppm"]), "sulfate": fnum(r["Sulfate_ppm"]), "nitrate": fnum(r["Nitrate_ppm"])}
    TS_DATE_CONFLICTS.extend(date_conflicts)
    extra = [
        ctx("comment", string_value(f"Source: {DRIVE_LABEL} 2024_SSO_Pilot Time Series/Chakraborty_Anions/20241220_IC_results_SSO_timeseries.xlsx; samples 0.2 um syringe filtered and stored at 4 C; single injection per sample")),
        ctx("comment", string_value("Samples are mapped by the workbook timepoint label (T1 = 2024-09-09, T2 = 2024-09-18); the workbook Date Sampled column disagrees with that label for " + ", ".join(c.split(" ")[0] for c in date_conflicts))),
    ]
    return build_ic(samples, table, 1, "chakraborty_ic_anions_sso_pilot_time_series_2024", "Chakraborty Lab ion chromatography anions (chloride, sulfate, nitrate) for the 2024 SSO pilot time series groundwater samples", extra), samples


def build_pump_toc() -> tuple[dict, list[str]]:
    data = rows(TOC_PUMP)
    samples: list[str] = []
    table: dict[tuple[str, int], tuple[float | None, float | None]] = {}
    for r in data:
        s = pump_sample(r["Well"], r["Zone"], r["sampling_period"], "ICTOC")
        if s not in samples:
            samples.append(s)
        table[(s, int(r["replicate"]))] = (fnum(r["TOC_ppm"]), fnum(r["TN_ppm"]))
    toc: list[float | None] = []
    toc_below: list[bool | None] = []
    tn: list[float | None] = []
    for s in samples:
        for rep in (1, 2, 3):
            c, n = table.get((s, rep), (None, None))
            if c is None:
                toc.append(None); toc_below.append(None)
            elif c <= 0:
                toc.append(None); toc_below.append(True)
            else:
                toc.append(round(c, 4)); toc_below.append(False)
            tn.append(None if n is None else round(n, 4))
    toc_ctx = [ctx("molecule from list", oterm_value("elemental carbon")), ctx("organic", boolean_value(True))]
    tn_ctx = [ctx("molecule from list", oterm_value("nitrogen atom"))]
    array_context = [
        ctx("instrument", string_value("Lotix combustion TOC-TN analyzer")),
        ctx("method", string_value("Catalytic combustion (680 to 1000 C) with NDIR detection for total organic carbon and chemiluminescence detection for total nitrogen; 7 mL sample diluted to 40 mL with nanopure water; triplicate measurements")),
        ctx("personnel", oterm_value("Romy Chakraborty")),
        ctx("comment", string_value(f"Source: {DRIVE_LABEL} 2024_SSO_Pump Test/Chakraborty_IC／TOC/20240416_TOC_TN_SSO_pumptest_corrected.xlsx")),
    ]
    obj = brick("chakraborty_toc_tn_sso_pump_test_2024", "Chakraborty Lab total organic carbon and total nitrogen for the 2024 SSO pump test groundwater samples", array_context, [sample_dim(samples), replicate_dim(3)], [
        float_var("concentration", "ppm", toc_ctx, toc),
        bool_var("below", below_ctx([ctx("molecule from list", oterm_value("elemental carbon")), ctx("organic", boolean_value(True))]), toc_below),
        float_var("concentration", "ppm", tn_ctx, tn),
    ])
    return obj, samples


def build_pump_isotopes() -> tuple[dict, list[str]]:
    summary = rows(ISO_PUMP, "Summary")
    nnum_to_sample = {}
    for r in summary:
        if r.get("N Number") and r.get("SampleID"):
            nnum_to_sample[str(int(r["N Number"]))] = str(r["SampleID"])
    wb = openpyxl.load_workbook(ISO_PUMP, read_only=True, data_only=True)
    ws = wb["Isotope Data"]
    it = ws.iter_rows(values_only=True)
    header = next(it)
    # columns: [None, None, N Number, Peak Ht, Calc NO3 uM (diluted), Dilution Factor, Est NO3 uM, d15N, d18O]
    per: dict[str, list[tuple]] = OrderedDict()
    for r in it:
        if r[2] is None:
            continue
        n = str(int(r[2]))
        per.setdefault(n, []).append((fnum(r[6]), fnum(r[7]), fnum(r[8]), fnum(r[5])))
    samples: list[str] = []
    no3: list[float | None] = []
    d15: list[float | None] = []
    d18: list[float | None] = []
    for n, reps in per.items():
        sid = nnum_to_sample[n]
        well, zone, period = sid.split("_")
        s = pump_sample(well, zone, period, "ISOTOPES")
        samples.append(s)
        if len(reps) != 2:
            raise ValueError(f"expected 2 analytical replicates for {sid}, found {len(reps)}")
        for c, a, b, _dil in reps:
            no3.append(None if c is None else round(c, 3))
            d15.append(None if a is None else round(a, 3))
            d18.append(None if b is None else round(b, 3))
    nitrate = [ctx("molecule from list", oterm_value("nitrate"))]
    array_context = [
        ctx("method", string_value("Denitrifier method for nitrate d15N and d18O (Wankel lab analysis of samples prepared by the Stahl lab); nitrate concentration estimated from the N2O peak height after dilution correction; two analytical replicates per sample")),
        ctx("personnel", oterm_value("Kristopher Hunt")),
        ctx("comment", string_value(f"Source: {DRIVE_LABEL} 2024_SSO_Pump Test/Stahl_Isotopes/20240416_TOC_TN_IC_ISOTOPES_SSO_pumptest.xlsx, sheet Isotope Data")),
        ctx("comment", string_value("The 200 L sample of the M5-SZ3 pumping series was recorded under the 100 L name in the field workbook; it is imported as M5-SZ3-20240329-ISOTOPES200")),
    ]
    obj = brick("stahl_nitrate_isotopes_sso_pump_test_2024", "Stahl Lab nitrate d15N and d18O isotope ratios and nitrate concentration for the 2024 SSO pump test groundwater samples", array_context, [sample_dim(samples), replicate_dim(2)], [
        float_var("concentration", "uM", nitrate, no3),
        float_var("isotope ratio", "per mil", nitrate + [ctx("isotope", oterm_value("nitrogen-15 atom"))], d15),
        float_var("isotope ratio", "per mil", nitrate + [ctx("isotope", oterm_value("oxygen-18 atom"))], d18),
    ])
    return obj, samples


# ----------------------------------------------------------------------------
# validation and packaging
# ----------------------------------------------------------------------------

def run_checkgeneric(json_path: Path) -> tuple[bool, str]:
    cp = [str(CHECKGENERIC_CLASSES), os.environ.get("CLASSPATH", "")] + [str(p) for p in JACKSON_JARS]
    proc = subprocess.run(["java", "-cp", os.pathsep.join(p for p in cp if p), "gov.lbl.enigma.app.CheckGeneric", str(json_path)], capture_output=True, text=True)
    out = (proc.stdout or "") + (proc.stderr or "")
    return proc.returncode == 0 and "Generic is OK!" in out, out


def check_shape(obj: dict) -> None:
    n = 1
    for d in obj["dim_context"]:
        n *= d["size"]
        for tv in d["typed_values"]:
            vals = tv["values"]
            arr = vals.get("object_refs") or vals.get("oterm_refs") or vals.get("int_values") or vals.get("string_values")
            if len(arr) != d["size"]:
                raise ValueError(f"{obj['name']}: dimension variable length {len(arr)} != {d['size']}")
    for v in obj["typed_values"]:
        arr = v["values"].get("float_values") or v["values"].get("boolean_values")
        if len(arr) != n:
            raise ValueError(f"{obj['name']}: variable length {len(arr)} != {n}")


def main() -> None:
    load_ontology_names()
    if OUT.exists():
        raise FileExistsError(f"refusing to overwrite {OUT}")
    for d in (JSON_OUT, STATIC_OUT, PROCESS_OUT, CHECK_OUT, REPORT_OUT):
        d.mkdir(parents=True)

    live_samples = {}
    with (EXPORT / "Sample.tsv").open() as h:
        for r in csv.DictReader(h, delimiter="\t"):
            live_samples[r["name"]] = r
    new_sample = "M5-SZ3-20240329-ISOTOPES200"
    template = live_samples["M5-SZ3-20240329-ISOTOPES100"]
    with (STATIC_OUT / f"Sample_additions_{RELEASE}.tsv").open("w", newline="") as h:
        w = csv.writer(h, delimiter="\t", lineterminator="\n")
        w.writerow(["name", "location", "depth", "elevation", "date", "time", "timezone", "material", "env_package", "description"])
        w.writerow([new_sample, template["location_id"], template["depth"], "null", "2024-03-29", "15:02", template["timezone"], template["material"], template["env_package"],
                    "Nitrate isotopes; 200 L pumped sample of the M5-SZ3 pumping series (recorded as ISOTOPES100 at 15:02 in the field metadata workbook)"])
    staged = set(live_samples) | {new_sample}

    builds = [
        ("ic_pump", build_pump_ic, "Romy Chakraborty", "2024-04-12"),
        ("toc_pump", build_pump_toc, "Romy Chakraborty", "2024-04-16"),
        ("ic_ts", build_ts_ic, "Romy Chakraborty", "2024-12-20"),
        ("iso_pump", build_pump_isotopes, "Kristopher Hunt", "2024-04-16"),
    ]
    results = []
    for key, fn, person, date in builds:
        obj, samples = fn()
        missing = [s for s in samples if s not in staged]
        if missing:
            raise ValueError(f"{obj['name']}: samples not in CORAL or staged additions: {missing}")
        check_shape(obj)
        path = JSON_OUT / obj["name"].replace(".hndarray", ".json")
        path.write_text(json.dumps(obj, indent=2, ensure_ascii=False) + "\n")
        ok, out = run_checkgeneric(path)
        (CHECK_OUT / (path.name + ".check")).write_text(out)
        if not ok:
            raise RuntimeError(f"CheckGeneric failed for {path}:\n{out[-2500:]}")
        results.append({"key": key, "name": obj["name"], "json": path.name, "samples": samples, "person": person, "date": date,
                        "shape": [d["size"] for d in obj["dim_context"]], "variables": len(obj["typed_values"]),
                        "new_sample_used": new_sample in samples})
        print(f"built {obj['name']} shape {[d['size'] for d in obj['dim_context']]} vars {len(obj['typed_values'])} CheckGeneric OK", flush=True)

    fieldnames = ["process", "person", "campaign", "protocol", "date_start", "date_end", "input_objects", "output_objects"]
    campaign = "Subsurface Observatory <ENIGMA:0000166>"
    sampling_path = PROCESS_OUT / f"process_sampling_sso_isotopes200_{RELEASE}.tsv"
    with sampling_path.open("w", newline="") as h:
        w = csv.DictWriter(h, fieldnames=fieldnames, delimiter="\t", lineterminator="\n")
        w.writeheader()
        w.writerow({"process": "Sampling <PROCESS:0000002>", "person": "Hazen Lab <ENIGMA:0000066>", "campaign": campaign, "protocol": "null",
                    "date_start": "2024-03-29", "date_end": "2024-03-29", "input_objects": f"Location: {template['location_id']}", "output_objects": f"Sample: {new_sample}"})
    assay_path = PROCESS_OUT / f"process_assay_geochemistry_sso_2024_{RELEASE}.tsv"
    with assay_path.open("w", newline="") as h:
        w = csv.DictWriter(h, fieldnames=fieldnames, delimiter="\t", lineterminator="\n")
        w.writeheader()
        for r in results:
            w.writerow({"process": "Assay Geochemistry <PROCESS:0000005>", "person": f"{r['person']} <{T[r['person']]}>", "campaign": campaign, "protocol": "null",
                        "date_start": r["date"], "date_end": r["date"], "input_objects": ", ".join(f"Sample: {s}" for s in r["samples"]), "output_objects": f"Generic: {r['name']}"})

    files = [f"static/{STATIC_OUT.name}/" if False else f"static/Sample_additions_{RELEASE}.tsv", f"process/{sampling_path.name}"] + [f"json/{r['json']}" for r in results] + [f"process/{assay_path.name}"]
    (OUT / "files_to_import.txt").write_text("\n".join(files) + "\n")
    script = [
        "# Paste into the CORAL import notebook after uploading every file in files_to_import.txt",
        "# (files are referenced by basename). Run the cells in order.",
        "",
        "# Cell 1: the missing 200 L isotope sample and its sampling process.",
        f"toolx.update_core('Sample_additions_{RELEASE}.tsv', 'Sample')",
        f"toolx.upload_process('Sampling', '{sampling_path.name}')",
        "",
        "# Cell 2: the four geochemistry bricks.",
    ] + [f"toolx.upload_brick('{r['json']}')" for r in results] + [
        "",
        "# Cell 3: the Assay Geochemistry processes linking samples to bricks.",
        f"toolx.upload_process('Assay Geochemistry', '{assay_path.name}')",
        "",
    ]
    (OUT / "import_to_coral.py").write_text("\n".join(script))
    summary = {"release": RELEASE, "bricks": results, "new_samples": [new_sample], "time_series_date_conflicts": TS_DATE_CONFLICTS, "source_export": str(EXPORT.relative_to(REPO))}
    (REPORT_OUT / "build_summary.json").write_text(json.dumps(summary, indent=2) + "\n")
    readme = [
        f"# 2024 SSO groundwater geochemistry CORAL import ({RELEASE})",
        "",
        "Four Chemical Measurement bricks from the Chakraborty and Stahl lab result workbooks on the team Google Drive, plus one static Sample addition.",
        "",
        "| Brick | Shape | Variables | Samples |",
        "| --- | --- | --- | --- |",
    ] + [f"| `{r['name']}` | {r['shape']} | {r['variables']} | {len(r['samples'])} |" for r in results] + [
        "",
        "* Ion chromatography concentrations are reported as the anion (nitrate as NO3-), 1:50 dilution corrected; zero readings are stored as null with the below-detection-limit flag set.",
        "* TOC/TN: total organic carbon (elemental carbon, organic=1) and total nitrogen (nitrogen atom), ppm, three technical replicates; zero TOC readings are null with the flag set.",
        "* Isotopes: nitrate d15N and d18O (per mil) and estimated nitrate (uM), two analytical replicates; blank replicate values are null.",
        f"* Pump-test PRE samples map to the morning `...ICTOC1` / `...ISOTOPES1` CORAL samples and POST to the afternoon `...2` samples; the M5-SZ3 0/100/200 L series maps to the 2024-03-29 samples. `{new_sample}` is added because the field workbook recorded the 15:02 200 L isotope sample under the 100 L name.",
        "* Time-series samples map by well, zone and timepoint label (T1 = 2024-09-09, T2 = 2024-09-18) to the `...ICTOC` samples; the workbook Date Sampled column contradicts the label on six rows (see `reports/build_summary.json`).",
        "",
        "Import order: `Sample_additions` and its Sampling process, then the bricks, then the Assay Geochemistry process file (see `import_to_coral.py`). All JSON files pass CheckGeneric (`check/`).",
        "",
    ]
    (OUT / "README.md").write_text("\n".join(readme))
    print(json.dumps({k: v for k, v in summary.items() if k != "bricks"}, indent=2))


if __name__ == "__main__":
    main()
