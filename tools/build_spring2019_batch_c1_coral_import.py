#!/usr/bin/env python3
"""Build the CORAL import package for batch C1 of IMPORT_PLAN_20261009 (Spring 2019).

Spring 2019 Integrated Sampling, nine wells (Area 1 background GW821E, GW822E,
GW823D; Area 2 MLSB5, MLSD5, MLSF5; Area 3 FW100-20, FW101-20, FW102-20),
five biweekly rounds from 2019-03-18 to 2019-05-15. CORAL already holds one
Sample per shipped bottle, named WELL-YYYY-MM-DD-FRACTION-R<assay>.

* Adams Lab ICP-MS metals: one HNO3-preserved bottle (assay 9.2, R92) per
  sampling; 52 elements, combined 1.2x/100x dilutions.
* Chakraborty Lab carbon and nitrogen: assay 13 (40 mL VOA, 0.2 um filtered,
  no headspace; R13, in round 5 two vials R131/R132); TN, TC, DOC, DIC for
  all wells, nitrate, nitrite, ammonium and DON for Areas 1 and 2.
* Hazen Lab AODC: assay 6 (10 mL formaldehyde-fixed, R6); processed and raw
  field counts.

Outputs: coral_import/spring2019_batch_c1_<release>/.
"""

from __future__ import annotations

import csv
import datetime as dt
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_sso_batch_ab_coral_import as ab  # shared term, brick and validation helpers

from build_sso_batch_ab_coral_import import (  # noqa: E402
    T, NAMES, ELEMENTS, term, oterm_value, ctx, sval, comment, typed, float_var, bool_var,
    sample_dim, molecule_dim, replicate_dim, below_ctx, brick, fnum, write_process, rel,
)

REPO = ab.REPO
DRIVE = ab.DRIVE
EXPORT = REPO / "sync-coral-to-berdl/exports/sync-20261016-152847/coral_export/static_tsv"
RELEASE = "20261016"
OUT = REPO / f"coral_import/spring2019_batch_c1_{RELEASE}"
SPRING = DRIVE / "2019_Spring_Integrated Sampling"
ADAMS = SPRING / "Adams Lab/2019_Groundwater_Time_Series_2HNO3_52_Final.xlsm"
ADAMS_README = SPRING / "Adams Lab/Readme.txt"
CN_A12 = SPRING / "Chakraborty Lab/C & N analysis_20190801.xlsx"
CN_A3 = SPRING / "Chakraborty Lab/Results_Area3.xlsx"
AODC = SPRING / "Hazen Lab/AODC/AODC_Active Sampling 2019_07112019_upload.xlsx"
CAMPAIGN = "Spring 2019 Integrated Sampling <ENIGMA:0000100>"
ONTOLOGY_UPDATE = REPO / "coral_import/ontology_update_20261016"
STAGED_ONT = Path("/tmp/claude-1000/-scratch-jmc-BERDL-ENIGMA-CORAL/e08931f7-7c7d-4424-a1ff-a5f1bad88ee0/scratchpad/stagedval")
AB_JSON = REPO / "coral_import/sso_batch_ab_20261009/json"


def stage_ontology() -> None:
    """Names and validation use the production ontologies plus the staged ME:0000506-0000512 update."""
    names = STAGED_ONT / "names"
    names.mkdir(parents=True, exist_ok=True)
    for f in ab.ONTOLOGIES.glob("*.obo"):
        link = names / f.name
        if f.name == "context_measurement_ontology.obo":
            continue
        if not link.exists():
            link.symlink_to(f)
    (names / "context_measurement_ontology.obo").write_bytes((ONTOLOGY_UPDATE / "production/context_measurement_ontology.obo").read_bytes())
    ab.ONTOLOGIES = names
    cp = (STAGED_ONT / "classpath.txt").read_text().strip()
    def run_staged(json_path):
        import subprocess
        proc = subprocess.run(["java", "-Xmx4g", "-cp", cp, "gov.lbl.enigma.app.CheckGenericStaged", str(json_path)], capture_output=True, text=True)
        out = (proc.stdout or "") + (proc.stderr or "")
        return proc.returncode == 0 and "Generic is OK!" in out, out
    ab.run_checkgeneric = run_staged

T.update({
    "elemental carbon": "CHEBI:33415", "nitrogen atom": "CHEBI:25555", "organic": "ME:0000346",
    "ppm": "UO:0000169", "Romy Chakraborty": "ENIGMA:0000054", "Spring 2019 Integrated Sampling": "ENIGMA:0000100",
})

WELLS = {"GW-821-E": "GW821E", "GW-822-E": "GW822E", "GW-823-D": "GW823D", "GW-823-E": "GW823D",
         "MLSB5": "MLSB5", "MLSD5": "MLSD5", "MLSF5": "MLSF5", "FW100-20": "FW100-20", "FW101-20": "FW101-20", "FW102-20": "FW102-20"}
WELL_ORDER = ["GW821E", "GW822E", "GW823D", "MLSB5", "MLSD5", "MLSF5", "FW100-20", "FW101-20", "FW102-20"]


def bottle(well: str, date: dt.date, fraction: str, rcode: str) -> str:
    return f"{WELLS[well]}-{date:%Y-%m-%d}-{fraction}-{rcode}"


def as_date(v) -> dt.date:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return dt.datetime.strptime(str(v), "%m%d%y").date()


def order_key(name: str):
    m = re.match(r"(.+?)-(\d{4}-\d{2}-\d{2})-", name)
    return (WELL_ORDER.index(m.group(1)), m.group(2))


# ----------------------------------------------------------------------------
# Adams Lab ICP-MS
# ----------------------------------------------------------------------------

def build_adams(staged: set[str]):
    wb = openpyxl.load_workbook(ADAMS, read_only=True, data_only=True)
    sheets = {}
    for kind, sn in (("conc", "Combined Conc (uM)"), ("sd", "Combined SD (uM)"), ("dl", "Combined DL Conc (uM)")):
        rows = list(wb[sn].iter_rows(values_only=True))
        symbols = [str(s).strip() for s in rows[2][7:59]]
        if symbols != [e[0] for e in ELEMENTS]:
            raise ValueError(f"{sn}: unexpected element order {symbols[:6]}")
        data = {}
        for r in rows[5:]:
            if not r[6]:
                continue
            orr = str(r[6]).strip()
            m = re.match(r"(.+)_(\d{6})_9\.2$", orr)
            if not m:
                raise ValueError(f"unexpected ORR sample id {orr}")
            name = bottle(m.group(1), as_date(m.group(2)), "UF", "R92")
            if as_date(r[1]) != as_date(m.group(2)):
                raise ValueError(f"{orr}: Date_Sampled {r[1]} disagrees with the ID")
            data[name] = (orr, [fnum(v) for v in r[7:59]])
        sheets[kind] = data
    samples = sorted(sheets["conc"], key=order_key)
    if any(set(sheets[k]) != set(samples) for k in sheets):
        raise ValueError("Adams sheets list different samples")
    conc = [v for s in samples for v in sheets["conc"][s][1]]
    dl = [v for s in samples for v in sheets["dl"][s][1]]
    below = [None if c is None or d is None else c < d for c, d in zip(conc, dl)]
    variables = [
        float_var("concentration", "uM", [], conc),
        float_var("standard deviation", "uM", [ctx("measurement", oterm_value("concentration"))], [v for s in samples for v in sheets["sd"][s][1]]),
        float_var("detection limit", "uM", [], dl),
        bool_var("below", below_ctx(), below),
    ]
    readme = ADAMS_README.read_text(errors="ignore").strip()
    array_context = [
        ctx("instrument", sval("Agilent 7900 ICP-MS (Adams Lab, University of Georgia)")),
        ctx("method", sval("ICP-MS of unfiltered groundwater preserved in HNO3 in the field (assay 9.2, 10 mL); each element is reported from the 1.2x or the 100x dilution run (Combined sheets); " + readme)),
        ctx("personnel", oterm_value("Farris Poole")),
        ctx("ENIGMA campaign", oterm_value("Spring 2019 Integrated Sampling")),
        comment("Variables: concentration, standard deviation of the concentration and detection limit (all micromolar, back-calculated to the original undiluted sample through each element's dilution factor), and below relative to the detection limit (true when the concentration is less than the detection limit; the lab marked these values red). Background equivalent concentrations are in the source workbook and are not loaded"),
        comment("Linked to the HNO3 bottle (R92) of each sampling; the in-field incubation (R91) and HCl (R93) bottles registered for the Adams Lab have no results in this workbook"),
        comment(f"Source: {rel(ADAMS)}, sheets Combined Conc/SD/DL Conc (uM); the ORR sample ID WELL_MMDDYY_9.2 gives the well, sampling date and bottle"),
    ]
    obj = brick("adams_metals_spring2019_integrated_sampling", "Adams Lab ICP-MS metal concentrations (52 elements) for the Spring 2019 Integrated Sampling groundwater samples",
                "chemical measurement", array_context, [sample_dim(samples), molecule_dim([f"element {sym}" for sym, _ in ELEMENTS])], variables)
    return obj, samples, {s: sheets["conc"][s][0] for s in samples}


# ----------------------------------------------------------------------------
# Chakraborty Lab carbon and nitrogen
# ----------------------------------------------------------------------------
NO3_N = 14.0067 / 62.0049
NH4_PER_N = 18.0385 / 14.0067


def build_cn(staged: set[str]):
    recs = {}
    def add(well, date, vals):
        rc = "R131" if date >= dt.date(2019, 5, 13) else "R13"
        name = bottle(well, date, "F02", rc)
        if name in recs:
            raise ValueError(f"duplicate C/N row {name}")
        recs[name] = {"well": well, **vals}
    rows = list(openpyxl.load_workbook(CN_A12, read_only=True, data_only=True)["Sheet2"].iter_rows(values_only=True))
    for r in rows[2:]:
        if not r or not r[0] or not r[1]:
            continue
        add(str(r[0]).strip(), as_date(r[1]), {"TN": fnum(r[3]), "NO3": fnum(r[4]), "NO2": fnum(r[5]), "NH4N": fnum(r[6]), "DON": fnum(r[7]),
                                                "TC": fnum(r[9]), "DOC": fnum(r[10]), "DIC": fnum(r[11])})
    rows = list(openpyxl.load_workbook(CN_A3, read_only=True, data_only=True)["Sheet1"].iter_rows(values_only=True))
    for r in rows[2:]:
        if not r or not r[0] or not r[1]:
            continue
        add(str(r[0]).strip(), as_date(r[1]), {"TN": fnum(r[3]), "TC": fnum(r[4]), "DIC": fnum(r[5]), "DOC": fnum(r[6])})
    samples = sorted(recs, key=order_key)
    def col(k):
        return [recs[s].get(k) for s in samples]
    def inorganic(k, scale=1.0):
        vals, bel = [], []
        for v in col(k):
            if v is None:
                vals.append(None); bel.append(None)
            elif v <= 0:
                vals.append(None); bel.append(True)
            else:
                vals.append(round(v * scale, 5)); bel.append(False)
        return vals, bel
    no3, no3_b = inorganic("NO3")
    no2, no2_b = inorganic("NO2")
    nh4, nh4_b = inorganic("NH4N", NH4_PER_N)
    org = lambda flag: [ctx("physiochemical state", sval("dissolved")), ctx("organic", {"scalar_type": "boolean", "boolean_value": 1 if flag else 0})]
    C, N = [ctx("molecule from list", oterm_value("elemental carbon"))], [ctx("molecule from list", oterm_value("nitrogen atom"))]
    mol = lambda k: [ctx("molecule from list", oterm_value(k))]
    variables = [
        float_var("concentration", "ppm", C, col("TC")),
        float_var("concentration", "ppm", C + org(True), col("DOC")),
        float_var("concentration", "ppm", C + org(False), col("DIC")),
        float_var("concentration", "ppm", N, col("TN")),
        float_var("concentration", "ppm", N + org(True), col("DON")),
        float_var("concentration", "ppm", mol("nitrate"), no3), bool_var("below", below_ctx(mol("nitrate")), no3_b),
        float_var("concentration", "ppm", mol("nitrite"), no2), bool_var("below", below_ctx(mol("nitrite")), no2_b),
        float_var("concentration", "ppm", mol("ammonium"), nh4), bool_var("below", below_ctx(mol("ammonium")), nh4_b),
    ]
    array_context = [
        ctx("personnel", oterm_value("Romy Chakraborty")),
        ctx("ENIGMA campaign", oterm_value("Spring 2019 Integrated Sampling")),
        ctx("method", sval("0.2 um filtered groundwater (assay 13, 40 mL VOA vial, no headspace, 4 C). Total carbon and total nitrogen by combustion analyzer; dissolved organic carbon measured and dissolved inorganic carbon = TC - DOC; nitrate and nitrite by ion chromatography; ammonium colorimetric; dissolved organic nitrogen = TN - nitrate-N - nitrite-N - ammonium-N (lab calculation)")),
        comment("Variables: total carbon (elemental carbon), dissolved organic carbon (state dissolved, organic true), dissolved inorganic carbon (state dissolved, organic false), total nitrogen (nitrogen atom), dissolved organic nitrogen (nitrogen atom, state dissolved, organic true), nitrate and nitrite (mg/L of the ion, as reported; the lab's detection-limit sheet converts them to N with 14.0067/62.0049 and 14.0067/46.005), and ammonium (reported as NH4-N, converted here to mg/L of ammonium ion by 18.0385/14.0067); all ppm = mg/L"),
        comment("Nitrate, nitrite and ammonium readings of 0 are stored as null with below relative to the detection limit set; detection limits (Results_Area3.xlsx, sheet Detection limit): TN 0.02 (quantitation 0.05), TC 0.5, TOC 0.1, nitrate 0.05 (0.2), nitrite 0.05 (0.2), NH4-N 0.04 mg/L. Area 3 wells (FW100-20, FW101-20, FW102-20) have TN, TC, DOC and DIC only"),
        comment("Round 5 (2019-05-13 to 05-15) was shipped as two VOA vials per well (R131, R132); results are linked to R131 and the lab did not record which vial each analysis used. The source sheet calls GW823D 'GW-823-E' on 2019-04-01 and 04-15; every other record names that well GW-823-D"),
        comment(f"Source: {rel(CN_A12)} (Areas 1 and 2) and {rel(CN_A3)} (Area 3)"),
    ]
    obj = brick("chakraborty_c_n_spring2019_integrated_sampling", "Chakraborty Lab carbon and nitrogen species (TC, DOC, DIC, TN, DON, nitrate, nitrite, ammonium) for the Spring 2019 Integrated Sampling groundwater samples",
                "chemical measurement", array_context, [sample_dim(samples)], variables)
    return obj, samples


# ----------------------------------------------------------------------------
# Hazen Lab AODC
# ----------------------------------------------------------------------------

def build_aodc():
    recs, blank = ab.read_aodc(AODC, iwc=False)
    samples = []
    for r in recs:
        m = re.match(r"(.+)_(\d{6})_6$", r["label"])
        if not m:
            raise ValueError(f"unexpected AODC id {r['label']}")
        samples.append(bottle(m.group(1), as_date(m.group(2)), "UF", "R6"))
    order = sorted(range(len(samples)), key=lambda i: order_key(samples[i]))
    recs = [recs[i] for i in order]; samples = [samples[i] for i in order]
    proc, raw = ab.build_aodc("spring2019_integrated_sampling", "Hazen Lab AODC processed data for the Spring 2019 Integrated Sampling groundwater samples", AODC, recs, samples, blank,
                              ["Assay 6: 10 mL unfiltered groundwater fixed with formaldehyde in the field (R6 bottles); 10 mL filtered per slide",
                               "This 2019-07-11 upload contains all 45 samplings; the earlier upload and the final-timepoint file hold subsets with identical values"])
    for o in (proc, raw):
        o["array_context"] = [c for c in o["array_context"] if c["value_type"]["oterm_ref"] != T["ENIGMA campaign"]] + [ctx("ENIGMA campaign", oterm_value("Spring 2019 Integrated Sampling"))]
        for c in o["array_context"]:
            if c["value_type"]["oterm_ref"] == T["method"]:
                old = c["value"]["string_value"]
                new = old.replace("1 mL of sample membrane filtered", "10 mL of sample membrane filtered").replace("/ volume filtered", "/ volume filtered (10 mL)")
                if new == old:
                    raise ValueError("AODC method text not updated")
                c["value"]["string_value"] = new
    return proc, raw, samples



STAT_NOTE = "statistic variables now name the summarized quantity (standard deviation; measurement = concentration)"
AODC_NOTE = "coefficient of variation and average now name the summarized quantity (measurement = count, category = cell); values unchanged"
MET_NOTE = ("values are now MS1 feature area (dimensionless) and controls carry control = negative control (extraction blanks) or positive control (QC standard); features now use m/z (dimensionless), retention time, CHEBI molecule from list and InChIKey/molecular formula as the reference compound, "
            "and spectral library compound name, match ID, match score and adduct for the GNPS match; 8 GNPS names lost their literal wrapping quotes; intensities unchanged")


def ab_v2_bricks(staged):
    """Rebuild the batch A/B bricks whose model changed, with the arguments the batch A/B build used."""
    out = []
    meta = ab.read_iwc_meta()
    p = ab.read_adams(ab.ADAMS_PUMP, "uM")
    obj, smp = ab.build_metals("adams_metals_sso_pump_test_2024", "Adams Lab ICP-MS metal concentrations (52 elements, sample minus blank) for the 2024 SSO pump test groundwater samples",
                               ab.ADAMS_PUMP, p, ab.pump_metals_sample, "uM", "2024-06-29",
                               ["Pre and Post samples of each port map to the start-of-pumping (...METALS1) and end-of-pumping (...METALS2) CORAL samples; the M5-SZ3 0, 100 and 200 L samples map to the 2024-03-29 METALS, METALS100 and METALS200 samples",
                                "Na, Mg, Al, K, Ca, Mn, Fe, Sr and U were measured at 375x dilution, the other elements at 1.2x"])
    out.append(("adams_metals_sso_pump_test_2024.hndarray", obj, smp, STAT_NOTE))
    p = ab.read_adams(ab.ADAMS_TS, "uM")
    obj, smp = ab.build_metals("adams_metals_sso_pilot_time_series_2024", "Adams Lab ICP-MS metal concentrations (52 elements, sample minus blank) for the 2024 SSO pilot time series groundwater samples (2024-09-09 and 2024-09-18)",
                               ab.ADAMS_TS, p, ab.ts_metals_sample_factory(p), "uM", "2024-10-09 (re-run 2025-09-17)",
                               ["Final data set of 2025-09-17: Na, Mg, Al, K, Ca, Mn, Sr and U re-measured at 600x dilution after the ICP-MS repair; Fe was not re-run and remains the 1.2x measurement, which the lab flagged as semi-quantitative",
                                "Data rows are joined to samples by tube number (HNO3-3000 to HNO3-3037, written HN03 in the data sheet) through the Sample Info sheet"])
    out.append(("adams_metals_sso_pilot_time_series_2024.hndarray", obj, smp, STAT_NOTE))
    p = ab.read_adams(ab.ADAMS_IWC, "uM")
    obj, smp = ab.build_metals("adams_metals_sso_initial_well_characterization_2023", "Adams Lab ICP-MS metal concentrations (52 elements, sample minus blank) for the 2023 SSO initial well characterization groundwater samples (28 ports)",
                               ab.ADAMS_IWC, p, ab.iwc_metals_sample_factory(meta), "uM", "2024-06-29",
                               ["Samples arrived unacidified despite the tube labels and were run as raw groundwater after about 3 months at 4 C (2.4x dilution; Na, Mg, Al, K, Ca, Mn, Fe, Sr and U at 750x)",
                                "Adams Lab IDs SSO-1 to SSO-28 follow the order of the Hazen metadata sheet (L7, L8, L9, M4, M5, M6, U1, U2, U3; SZ2, SZ1, VSZ within each well)"])
    out.append(("adams_metals_sso_initial_well_characterization_2023.hndarray", obj, smp, STAT_NOTE))
    p = ab.read_adams(ab.ADAMS_SED, "mg per Kg")
    obj, smp = ab.build_metals("adams_metals_sso_sediment_cores_2023", "Adams Lab ICP-MS metal concentrations (52 elements, mg per kg dry sediment, sample minus blank) for the 2023 SSO sediment core samples",
                               ab.ADAMS_SED, p, ab.sediment_sample, "mg/kg", "2024-06-29",
                               ["Concentrations are per kg dry sediment from the EPA 3051A microwave nitric-acid digest (the sample weight is the dry mass digested and the sample volume is the digest resuspension volume; 250x final dilution, Mg, Al, K, Ca, Mn and Fe at 25,000x); detection limits and background equivalent concentrations are computed for a nominal 0.5 g",
                                "Sample IDs SSO-<well>-C<n> are the CORAL -00 subsamples; the Adams workbook calls the M4 cores SSO-M4 (CORAL SSO-M4A) and notes that some hand-written IDs were matched approximately"],
                               sample_extra=ab.sediment_sample_extra)
    out.append(("adams_metals_sso_sediment_cores_2023.hndarray", obj, smp, STAT_NOTE))
    recs, blank = ab.read_aodc(ab.AODC_PUMP, iwc=False)
    smp = [ab.pump_aodc_sample(r["label"]) for r in recs]
    proc, _ = ab.build_aodc("sso_pump_test_2024", "Hazen Lab AODC processed data for the 2024 SSO pump test groundwater samples", ab.AODC_PUMP, recs, smp, blank,
                            ["PRE and POST labels map to the start-of-pumping (...AODC1) and end-of-pumping (...AODC2) CORAL samples; M5-SZ3 0L, 100L and 200L map to the 2024-03-29 AODC, AODC100 and AODC200 samples; the workbook has no count dates"])
    out.append(("hazen_aodc_sso_pump_test_2024.hndarray", proc, smp, AODC_NOTE))
    recs, blank = ab.read_aodc(ab.AODC_IWC, iwc=True)
    smp = [meta[r["label"].split()[0].replace("-", "_")]["aodc"] for r in recs]
    proc, _ = ab.build_aodc("sso_initial_well_characterization_2023", "Hazen Lab AODC processed data for the 2023 SSO initial well characterization groundwater samples", ab.AODC_IWC, recs, smp, blank,
                            ["Each port's second workbook row (_INT, intracellular count) has no field tallies and is not loaded; the L7-SZ2 slide is annotated 'Clumping'"])
    out.append(("hazen_aodc_sso_initial_well_characterization_2023.hndarray", proc, smp, AODC_NOTE))
    for key, pattern, polarity, chrom, run_date in ab.RUNS:
        obj, smp, _ = ab.build_metabolomics(key, pattern, polarity, chrom, run_date, staged)
        out.append((f"northen_metabolomics_sso_sediment_2023_{key}.hndarray", obj, smp, MET_NOTE))
    return out


def _arrays(o):
    """Measured values and sample references, which a model change must not alter."""
    vals = [tv["values"].get("float_values") or tv["values"].get("int_values") or tv["values"].get("boolean_values") for tv in o["typed_values"]]
    refs = [tv["values"].get("object_refs") for d in o["dim_context"] for tv in d["typed_values"] if tv["values"].get("object_refs")]
    return vals, refs, [d["size"] for d in o["dim_context"]]


def check_same_values(v1, v2):
    if _arrays(v1) != _arrays(v2):
        raise ValueError(f"{v2['name']}: measured values, sample references or shape differ from {v1['name']}")

# ----------------------------------------------------------------------------

def main() -> None:
    stage_ontology()
    ab.load_ontology_names()
    if OUT.exists():
        raise FileExistsError(f"refusing to overwrite {OUT}")
    ab.OUT = OUT
    ab.JSON_OUT, ab.STATIC_OUT, ab.PROCESS_OUT, ab.CHECK_OUT, ab.REPORT_OUT = (OUT / d for d in ("json", "static", "process", "check", "reports"))
    for d in (ab.JSON_OUT, ab.STATIC_OUT, ab.PROCESS_OUT, ab.CHECK_OUT, ab.REPORT_OUT):
        d.mkdir(parents=True)
    live = {}
    with (EXPORT / "Sample.tsv").open() as h:
        for r in csv.DictReader(h, delimiter="\t"):
            live[r["name"]] = r
    summary = {"release": RELEASE, "source_export": str(EXPORT.relative_to(REPO)), "bricks": [], "new_samples": []}

    adams, adams_samples, adams_ids = build_adams(set(live))
    new = [s for s in adams_samples if s not in live]
    if new != ["MLSB5-2019-04-30-UF-R92"]:
        raise ValueError(f"unexpected missing Adams samples {new}")
    tmpl = live["MLSB5-2019-04-30-UF-R3"]
    sample_path = ab.STATIC_OUT / f"Sample_additions_{RELEASE}.tsv"
    with sample_path.open("w", newline="") as h:
        w = csv.writer(h, delimiter="\t", lineterminator="\n")
        w.writerow(["name", "location", "depth", "elevation", "date", "time", "timezone", "material", "env_package", "description"])
        w.writerow([new[0], tmpl["location_id"], tmpl["depth"], "null", tmpl["date"], "null", "null", tmpl["material"], tmpl["env_package"],
                    "Adams; HNO3-preserved metals bottle (assay 9.2). Not in the 2019 sample registry, which lists no Adams bottles for this sampling; the Adams Lab reports results for MLSB5_043019_9.2"])
    summary["new_samples"] = new
    staged = set(live) | set(new)
    sampling_path = ab.PROCESS_OUT / f"process_sampling_spring2019_{RELEASE}.tsv"
    write_process(sampling_path, [{"process": "Sampling <PROCESS:0000002>", "person": "Hazen Lab <ENIGMA:0000066>", "campaign": CAMPAIGN, "protocol": "null",
                                   "date_start": "2019-04-30", "date_end": "2019-04-30", "input_objects": f"Location: {tmpl['location_id']}", "output_objects": f"Sample: {new[0]}"}])

    processes = defaultdict(list)
    def record(obj, samples, process, person, d0, d1, protocol="null", outputs=None, extra_inputs=()):
        missing = [s for s in samples if s not in staged]
        if missing:
            raise ValueError(f"{obj['name']}: samples not staged: {missing[:5]}")
        path = ab.write_brick(obj)
        summary["bricks"].append({"name": obj["name"], "json": path.name, "shape": [d["size"] for d in obj["dim_context"]], "variables": len(obj["typed_values"]), "samples": len(samples)})
        if process:
            processes[process].append({"process": process, "person": person, "campaign": CAMPAIGN, "protocol": protocol, "date_start": d0, "date_end": d1,
                                       "input_objects": ", ".join(f"Sample: {s}" for s in list(samples) + list(extra_inputs)), "output_objects": outputs or f"Generic: {obj['name']}"})

    GEO, CELLS = "Assay Geochemistry <PROCESS:0000005>", "Assay Cell Counts <PROCESS:0000008>"
    record(adams, adams_samples, GEO, f"Farris Poole <{T['Farris Poole']}>", "2019-05-30", "2019-06-19")
    cn, cn_samples = build_cn(staged)
    r132 = [s.replace("-R131", "-R132") for s in cn_samples if s.endswith("-R131")]
    missing132 = [s for s in r132 if s not in staged]
    if missing132:
        raise ValueError(f"R132 vials missing {missing132}")
    record(cn, cn_samples, GEO, f"Romy Chakraborty <{T['Romy Chakraborty']}>", "2019-04-17", "2019-08-01", extra_inputs=r132)
    proc, raw, aodc_samples = build_aodc()
    record(raw, aodc_samples, None, None, None, None)
    record(proc, aodc_samples, CELLS, f"Terry Hazen <{T['Terry Hazen']}>", "2019-03-18", "2019-07-11", protocol="hazen-2016-aodc",
           outputs=f"Generic: {proc['name']}, Generic: {raw['name']}")

    # Superseding versions of batch A/B bricks already in CORAL: statistics now name what they
    # summarize (standard deviation of concentration; CV and average of cell count), and the
    # metabolomics features use retention time, m/z, CHEBI and GNPS spectral-match microtypes.
    today = f"{RELEASE[:4]}-{RELEASE[4:6]}-{RELEASE[6:]}"
    for old, obj, msamples, note in ab_v2_bricks(staged):
        v1 = json.load(open(AB_JSON / old.replace(".hndarray", ".json")))
        obj["name"] = old.replace(".hndarray", "_v2.hndarray")
        if not any(c["value_type"]["oterm_ref"] == T["ENIGMA campaign"] for c in obj["array_context"]):
            obj["array_context"].append(ctx("ENIGMA campaign", oterm_value("Subsurface Observatory")))  # as on the assay processes
        obj["array_context"].append(comment(f"Version 2 of {old}: {note}"))
        check_same_values(v1, obj)
        record(obj, msamples, None, None, None, None)
        summary.setdefault("superseded", []).append({"old": old, "new": obj["name"], "change": note})
        processes["Update Data <PROCESS:0000053>"].append({"process": "Update Data <PROCESS:0000053>", "person": "John-Marc Chandonia <ENIGMA:0000057>",
            "campaign": "Subsurface Observatory <ENIGMA:0000166>", "protocol": "null", "date_start": today, "date_end": today,
            "input_objects": f"Generic: {old}", "output_objects": f"Generic: {obj['name']}"})

    process_files = []
    for process, rows in processes.items():
        slug = process.split(" <")[0].lower().replace(" ", "_")
        path = ab.PROCESS_OUT / f"process_{slug}_spring2019_{RELEASE}.tsv"
        write_process(path, rows)
        process_files.append((process.split(" <")[0], path.name))
    files = [f"static/{sample_path.name}", f"process/{sampling_path.name}"] + [f"json/{b['json']}" for b in summary["bricks"]] + [f"process/{n}" for _, n in process_files]
    (OUT / "files_to_import.txt").write_text("\n".join(files) + "\n")
    script = ["# Before this notebook: run coral_import/ontology_update_20261016/install_me0000506_0000512.sh (ME:0000506-ME:0000512).",
              "# Paste into the CORAL import notebook after uploading every file in files_to_import.txt", "",
              "# Cell 1: the one missing Adams bottle sample and its Sampling process.",
              f"toolx.update_core('{sample_path.name}', 'Sample')", f"toolx.upload_process('Sampling', '{sampling_path.name}')", "",
              "# Cell 2: the bricks."] + [f"toolx.upload_brick('{b['json']}')" for b in summary["bricks"]] + ["", "# Cell 3: the assay processes."] + \
             [f"toolx.upload_process('{p}', '{n}')" for p, n in process_files] + [""]
    (OUT / "import_to_coral.py").write_text("\n".join(script))
    summary["adams_ids"] = adams_ids
    (ab.REPORT_OUT / "build_summary.json").write_text(json.dumps(summary, indent=2, default=str) + "\n")
    readme = [f"# Spring 2019 batch C1 CORAL import ({RELEASE})", "",
              "Adams metals, Chakraborty carbon and nitrogen, and Hazen AODC for the Spring 2019 Integrated Sampling, built by `tools/build_spring2019_batch_c1_coral_import.py`.", "",
              "| Brick | Shape | Variables | Samples |", "| --- | --- | --- | --- |"] + \
             [f"| `{b['name']}` | {b['shape']} | {b['variables']} | {b['samples']} |" for b in summary["bricks"]] + ["",
              "* Bottles: Adams R92 (HNO3), Chakraborty R13 (R131 in round 5, R132 listed as a process input), Hazen R6.",
              "* Superseding `_v2` bricks (one Update Data process each) for ten batch A/B bricks already in CORAL: four Adams metals and two processed AODC bricks (statistics now name the summarized quantity) and the four sediment metabolomics feature tables (MS1 feature area, m/z, retention time, CHEBI and GNPS spectral-match microtypes and control types; ontology update ME:0000506-ME:0000512). Measured values are unchanged.",
              "* Requires `coral_import/ontology_update_20261016` to be installed before the bricks are loaded.",
              f"* Static addition: `{new[0]}` (the registry has no Adams bottles for that sampling, but the lab reports results for it).",
              "", "Import order: the sample and its Sampling process, then the bricks, then the assay processes (`import_to_coral.py`). All JSON files pass CheckGeneric (`check/`).", ""]
    (OUT / "README.md").write_text("\n".join(readme))
    print(json.dumps({k: v for k, v in summary.items() if k not in ("adams_ids",)}, indent=1, default=str))


if __name__ == "__main__":
    main()
